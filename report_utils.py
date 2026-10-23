"""Core data and export helpers for Business Report Automator."""
from __future__ import annotations
from io import BytesIO
from pathlib import Path
from copy import copy
from typing import Any
import pandas as pd

NONE_LABEL = "— Not mapped —"


def read_data_file(file_obj: Any, filename: str | None = None) -> pd.DataFrame:
    """Read CSV or Excel data, trim headers, and validate the basic structure."""
    name = filename or getattr(file_obj, "name", "upload.csv")
    suffix = Path(name).suffix.lower()
    if suffix == ".csv":
        try:
            frame = pd.read_csv(file_obj, encoding="utf-8-sig")
        except UnicodeDecodeError:
            if hasattr(file_obj, "seek"):
                file_obj.seek(0)
            frame = pd.read_csv(file_obj, encoding="latin-1")
    elif suffix in {".xlsx", ".xlsm"}:
        frame = pd.read_excel(file_obj, engine="openpyxl")
    elif suffix == ".xls":
        raise ValueError("Legacy .xls files are not supported. Save the workbook as .xlsx or .csv.")
    else:
        raise ValueError("Unsupported file type. Upload a CSV or Excel (.xlsx/.xlsm) file.")
    frame.columns = [str(c).strip() for c in frame.columns]
    if len(frame.columns) == 0:
        raise ValueError("The file has no columns.")
    if frame.empty:
        raise ValueError("The file has no data rows.")
    if frame.columns.duplicated().any():
        duplicate_headers = frame.columns[frame.columns.duplicated()].tolist()
        raise ValueError(f"Column names must be unique. Duplicate headers: {', '.join(duplicate_headers)}")
    return frame


def guess_column(columns: list[str], candidates: list[str]) -> str:
    """Suggest a likely column using common header names."""
    normalised = {str(col).strip().lower().replace("_", " "): col for col in columns}
    for candidate in candidates:
        key = candidate.lower().replace("_", " ")
        if key in normalised:
            return normalised[key]
    for col in columns:
        text = str(col).strip().lower().replace("_", " ")
        if any(candidate.lower().replace("_", " ") in text for candidate in candidates):
            return col
    return NONE_LABEL


def parse_numeric(series: pd.Series) -> pd.Series:
    """Parse typical spreadsheet numbers, including currency symbols and commas."""
    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series, errors="coerce")
    text = series.astype("string").str.strip()
    parenthetical_negative = text.str.match(r"^\(.*\)$", na=False)
    cleaned = text.str.replace(r"[₹$€£,\s]", "", regex=True)
    cleaned = cleaned.str.replace(r"^\((.*)\)$", r"-\1", regex=True)
    cleaned = cleaned.str.replace(r"[^0-9.\-+]", "", regex=True)
    parsed = pd.to_numeric(cleaned, errors="coerce")
    mask = parenthetical_negative & parsed.notna()
    parsed.loc[mask] = -parsed.loc[mask].abs()
    return parsed


def prepare_analysis(
    df_raw: pd.DataFrame,
    *,
    sales_col: str = NONE_LABEL,
    date_col: str = NONE_LABEL,
    category_col: str = NONE_LABEL,
    product_col: str = NONE_LABEL,
    order_col: str = NONE_LABEL,
    quantity_col: str = NONE_LABEL,
    unit_price_col: str = NONE_LABEL,
    revenue_mode: str = "column",
    remove_duplicates: bool = True,
    trim_text: bool = True,
    dayfirst_dates: bool = False,
) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    """Return cleaned raw rows, an analysis frame, and data-quality summary."""
    source = df_raw.copy()
    if trim_text:
        for col in source.select_dtypes(include=["object", "string"]).columns:
            source[col] = source[col].map(lambda value: value.strip() if isinstance(value, str) else value)

    original_rows = len(source)
    duplicate_count = int(source.duplicated().sum())
    missing_by_column = source.isna().sum()
    missing_cells = int(missing_by_column.sum())
    rows_with_missing = int(source.isna().any(axis=1).sum())
    cleaned = source.drop_duplicates().copy() if remove_duplicates else source.copy()

    if revenue_mode == "quantity_price":
        if quantity_col == NONE_LABEL or unit_price_col == NONE_LABEL:
            raise ValueError("Choose both quantity and unit-price columns to calculate revenue.")
        sales = parse_numeric(cleaned[quantity_col]) * parse_numeric(cleaned[unit_price_col])
    else:
        if sales_col == NONE_LABEL:
            raise ValueError("Choose a sales/revenue column or calculate revenue from quantity × unit price.")
        sales = parse_numeric(cleaned[sales_col])

    analysis = cleaned.copy()
    analysis["_report_sales"] = sales
    if date_col != NONE_LABEL:
        analysis["_report_date"] = pd.to_datetime(analysis[date_col], errors="coerce", dayfirst=dayfirst_dates)
    else:
        analysis["_report_date"] = pd.NaT

    for internal, col in [("_report_category", category_col), ("_report_product", product_col), ("_report_order", order_col)]:
        if col != NONE_LABEL:
            analysis[internal] = analysis[col].fillna("Unspecified").astype(str).replace("", "Unspecified")
        else:
            analysis[internal] = "Unspecified"

    invalid_sales = int(analysis["_report_sales"].isna().sum())
    invalid_dates = int(analysis["_report_date"].isna().sum()) if date_col != NONE_LABEL else 0
    negative_sales = int((analysis["_report_sales"] < 0).sum())
    quality = {
        "original_rows": int(original_rows),
        "output_rows": int(len(cleaned)),
        "duplicate_rows_detected": duplicate_count,
        "duplicate_rows_removed": duplicate_count if remove_duplicates else 0,
        "rows_with_missing_values": rows_with_missing,
        "missing_cells": missing_cells,
        "invalid_sales_values": invalid_sales,
        "invalid_dates": invalid_dates,
        "negative_sales_rows": negative_sales,
        "missing_by_column": missing_by_column,
    }
    return cleaned, analysis, quality


def visible_data(frame: pd.DataFrame) -> pd.DataFrame:
    """Drop internal calculation columns before presenting or exporting rows."""
    return frame[[col for col in frame.columns if not str(col).startswith("_report_")]].copy()


def format_inr(value: float) -> str:
    return f"₹{value:,.2f}"


def build_excel_report(
    filtered_analysis: pd.DataFrame,
    *,
    source_name: str,
    quality: dict[str, Any],
    currency_symbol: str = "₹",
) -> bytes:
    """Build a multi-sheet Excel report for the current filtered view."""
    frame = filtered_analysis.copy()
    valid_sales = frame["_report_sales"].dropna()
    total = float(valid_sales.sum()) if not valid_sales.empty else 0.0
    average = float(valid_sales.mean()) if not valid_sales.empty else 0.0
    valid_count = int(frame["_report_sales"].notna().sum())

    overview = pd.DataFrame([
        ["Source file", source_name],
        ["Rows in current view", int(len(frame))],
        ["Rows with numeric sales", valid_count],
        ["Total sales", total],
        ["Average sales per row", average],
        ["Distinct categories", int(frame["_report_category"].nunique())],
        ["Distinct products", int(frame["_report_product"].nunique())],
        ["Report currency label", currency_symbol],
    ], columns=["Metric", "Value"])

    monthly = pd.DataFrame(columns=["Month", "Sales"])
    dated = frame.dropna(subset=["_report_date", "_report_sales"]).copy()
    if not dated.empty:
        dated["Month"] = dated["_report_date"].dt.to_period("M").astype(str)
        monthly = dated.groupby("Month", as_index=False)["_report_sales"].sum().rename(columns={"_report_sales": "Sales"})

    categories = frame.dropna(subset=["_report_sales"]).groupby("_report_category", as_index=False)["_report_sales"].sum()
    categories = categories.sort_values("_report_sales", ascending=False).rename(columns={"_report_category": "Category", "_report_sales": "Sales"})
    products = frame.dropna(subset=["_report_sales"]).groupby("_report_product", as_index=False)["_report_sales"].sum()
    products = products.sort_values("_report_sales", ascending=False).head(25).rename(columns={"_report_product": "Product", "_report_sales": "Sales"})

    quality_rows = [
        ["Original rows", quality["original_rows"]],
        ["Rows in cleaned source", quality["output_rows"]],
        ["Exact duplicate rows detected", quality["duplicate_rows_detected"]],
        ["Exact duplicate rows removed", quality["duplicate_rows_removed"]],
        ["Rows with missing values", quality["rows_with_missing_values"]],
        ["Missing cells", quality["missing_cells"]],
        ["Invalid sales values", quality["invalid_sales_values"]],
        ["Invalid/unreadable dates", quality["invalid_dates"]],
        ["Negative sales rows (may be refunds/adjustments)", quality["negative_sales_rows"]],
    ]
    quality_sheet = pd.DataFrame(quality_rows, columns=["Check", "Count"])
    missing_detail = quality["missing_by_column"].rename("Missing cells").rename_axis("Column").reset_index()

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        overview.to_excel(writer, sheet_name="Overview", index=False)
        monthly.to_excel(writer, sheet_name="Monthly trend", index=False)
        categories.to_excel(writer, sheet_name="Categories", index=False)
        products.to_excel(writer, sheet_name="Top products", index=False)
        quality_sheet.to_excel(writer, sheet_name="Data quality", index=False, startrow=0)
        missing_detail.to_excel(writer, sheet_name="Data quality", index=False, startrow=len(quality_sheet) + 3)
        visible_data(frame).to_excel(writer, sheet_name="Filtered data", index=False)
        from openpyxl.styles import PatternFill
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cells in ws.columns:
                letter = cells[0].column_letter
                max_length = max((len(str(cell.value)) if cell.value is not None else 0) for cell in cells[:1000])
                ws.column_dimensions[letter].width = min(max(max_length + 2, 12), 36)
            for cell in ws[1]:
                font = copy(cell.font)
                font.bold = True
                font.color = "FFFFFF"
                cell.font = font
                cell.fill = PatternFill("solid", fgColor="243B53")
    return buffer.getvalue()
